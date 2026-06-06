#!/usr/bin/env python3
"""
ClimateLenz Align-Paste Tool
A web application that takes copied material and templates (Word/Excel/PPT)
and produces aligned, formatted output that can be downloaded.

Author: ClimateLenz
"""

from flask import Flask, render_template, request, send_file, jsonify, flash, redirect
from werkzeug.utils import secure_filename
import os
import io
import re
from datetime import datetime
import zipfile

# Document processing libraries
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches
    from pptx.enum.text import PP_ALIGN
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

app = Flask(__name__)
app.secret_key = 'climatelenz-secret-key-2026'
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max upload

# Configuration
UPLOAD_FOLDER = 'uploads'
OUTPUT_FOLDER = 'outputs'
ALLOWED_EXTENSIONS = {'docx', 'xlsx', 'pptx', 'txt', 'md'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['OUTPUT_FOLDER'] = OUTPUT_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


class MaterialProcessor:
    """Processes copied material and aligns it with templates."""
    
    def __init__(self, material_text):
        self.material = material_text
        self.sections = self._parse_material()
    
    def _parse_material(self):
        """Parse material into structured sections."""
        sections = []
        
        # Split by common section headers (numbered, bold, or markdown headers)
        lines = self.material.split('\n')
        current_section = {'title': 'Introduction', 'content': []}
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Detect section headers
            is_header = False
            
            # Pattern 1: Numbered headers (1. Title, 1.1 Title, etc.)
            if re.match(r'^\d+(\.\d+)*\s+[A-Z]', line):
                is_header = True
            # Pattern 2: Markdown headers
            elif line.startswith('#') or line.startswith('**') and line.endswith('**'):
                is_header = True
            # Pattern 3: ALL CAPS short lines
            elif line.isupper() and len(line) < 100:
                is_header = True
            # Pattern 4: Lines ending with colon
            elif line.endswith(':') and len(line) < 100:
                is_header = True
            
            if is_header and current_section['content']:
                sections.append(current_section)
                current_section = {'title': line, 'content': []}
            else:
                current_section['content'].append(line)
        
        if current_section['content']:
            sections.append(current_section)
        
        return sections
    
    def get_formatted_text(self):
        """Return material as formatted text."""
        return self.material
    
    def get_section_titles(self):
        """Return list of section titles."""
        return [s['title'] for s in self.sections]


class TemplateAligner:
    """Aligns material into document templates."""
    
    def __init__(self, material_processor):
        self.material = material_processor
    
    def align_to_word(self, template_path=None):
        """Align material to Word template or create new document."""
        if not DOCX_AVAILABLE:
            return None
        
        try:
            if template_path and os.path.exists(template_path):
                doc = Document(template_path)
            else:
                doc = Document()
        except:
            doc = Document()
        
        # Add title
        title = doc.add_heading('ClimateLenz Aligned Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add date
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_run = date_para.add_run(f'Generated: {datetime.now().strftime("%B %d, %Y")}')
        date_run.font.size = Pt(10)
        date_run.font.italic = True
        
        doc.add_paragraph()  # Spacing
        
        # Add sections
        for section in self.material.sections:
            # Section heading
            heading = doc.add_heading(section['title'], level=2)
            
            # Section content
            for content_line in section['content']:
                para = doc.add_paragraph()
                run = para.add_run(content_line)
                run.font.size = Pt(11)
                run.font.name = 'Calibri'
                
                # Detect and format lists
                if content_line.startswith('•') or content_line.startswith('-') or content_line.startswith('*'):
                    para.style = 'List Bullet'
                elif re.match(r'^\d+\.', content_line):
                    para.style = 'List Number'
        
        # Add footer
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run('— Prepared by ClimateLenz —')
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = RGBColor(128, 128, 128)
        
        return doc
    
    def align_to_excel(self, template_path=None):
        """Align material to Excel template or create new workbook."""
        if not EXCEL_AVAILABLE:
            return None
        
        try:
            if template_path and os.path.exists(template_path):
                wb = load_workbook(template_path)
                ws = wb.active
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "Aligned Content"
        except:
            wb = Workbook()
            ws = wb.active
            ws.title = "Aligned Content"
        
        # Styles
        header_font = Font(bold=True, size=14, color="FFFFFF")
        header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        content_font = Font(size=11)
        content_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Clear existing content if new workbook
        if not template_path:
            ws.delete_rows(1, ws.max_row)
        
        # Add title row
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = "ClimateLenz - Aligned Material Report"
        title_cell.font = Font(bold=True, size=16, color="2E7D32")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30
        
        # Add metadata
        ws.merge_cells('A2:D2')
        meta_cell = ws['A2']
        meta_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Sections: {len(self.material.sections)}"
        meta_cell.font = Font(italic=True, size=10, color="666666")
        meta_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Headers
        headers = ['Section', 'Title', 'Content', 'Status']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        ws.row_dimensions[4].height = 25
        
        # Content rows
        row = 5
        for idx, section in enumerate(self.material.sections, 1):
            content_text = '\n'.join(section['content'])
            
            data = [
                idx,
                section['title'],
                content_text,
                "Aligned"
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = value
                cell.font = content_font
                cell.alignment = content_alignment
                cell.border = thin_border
            
            ws.row_dimensions[row].height = max(30, len(content_text) // 50 * 15)
            row += 1
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 60
        ws.column_dimensions['D'].width = 15
        
        # Add summary sheet
        summary = wb.create_sheet("Summary")
        summary['A1'] = "Total Sections"
        summary['B1'] = len(self.material.sections)
        summary['A2'] = "Total Content Lines"
        summary['B2'] = sum(len(s['content']) for s in self.material.sections)
        summary['A3'] = "Generated By"
        summary['B3'] = "ClimateLenz Align-Paste Tool"
        
        for cell in ['A1', 'A2', 'A3']:
            summary[cell].font = Font(bold=True)
        
        return wb
    
    def align_to_pptx(self, template_path=None):
        """Align material to PowerPoint template or create new presentation."""
        if not PPTX_AVAILABLE:
            return None
        
        try:
            if template_path and os.path.exists(template_path):
                prs = Presentation(template_path)
            else:
                prs = Presentation()
        except:
            prs = Presentation()
        
        # Title slide
        title_slide_layout = prs.slide_layouts[0]  # Title slide
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = "ClimateLenz"
        subtitle.text = f"Aligned Material Report\n{datetime.now().strftime('%B %d, %Y')}"
        
        # Content slides - one per section
        for section in self.material.sections:
            bullet_slide_layout = prs.slide_layouts[1]  # Title and Content
            slide = prs.slides.add_slide(bullet_slide_layout)
            
            shapes = slide.shapes
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]
            
            title_shape.text = section['title']
            
            tf = body_shape.text_frame
            tf.text = section['content'][0] if section['content'] else ""
            
            for line in section['content'][1:]:
                p = tf.add_paragraph()
                p.text = line
                p.level = 0
        
        # Summary slide
        summary_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(summary_slide_layout)
        slide.shapes.title.text = "Summary"
        
        tf = slide.shapes.placeholders[1].text_frame
        tf.text = f"Total Sections: {len(self.material.sections)}"
        p = tf.add_paragraph()
        p.text = f"Generated by ClimateLenz Align-Paste Tool"
        p.level = 0
        
        return prs


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/process', methods=['POST'])
def process_material():
    """Process copied material and template, return aligned output."""
    
    # Get material text
    material_text = request.form.get('material', '').strip()
    if not material_text:
        return jsonify({'error': 'No material provided'}), 400
    
    # Get alignment options
    output_format = request.form.get('format', 'docx')
    align_mode = request.form.get('align_mode', 'auto')
    
    # Process material
    processor = MaterialProcessor(material_text)
    aligner = TemplateAligner(processor)
    
    # Handle template file
    template_path = None
    if 'template' in request.files:
        file = request.files['template']
        if file and file.filename and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            template_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(template_path)
    
    # Generate output
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_filename = f"climatelenz_aligned_{timestamp}"
    
    try:
        if output_format == 'docx':
            doc = aligner.align_to_word(template_path)
            if doc:
                output_path = os.path.join(app.config['OUTPUT_FOLDER'], f"{output_filename}.docx")
                doc.save(output_path)
                return send_file(output_path, as_attachment=True, 
                               download_name=f"ClimateLenz_Report_{timestamp}.docx")
            else:
                return jsonify({'error': 'python-docx not installed'}), 500
        
        elif output_format == 'xlsx':
            wb = aligner.align_to_excel(template_path)
            if wb:
                output_path = os.path.join(app.config['OUTPUT_FOLDER'], f"{output_filename}.xlsx")
                wb.save(output_path)
                return send_file(output_path, as_attachment=True,
                               download_name=f"ClimateLenz_Report_{timestamp}.xlsx")
            else:
                return jsonify({'error': 'openpyxl not installed'}), 500
        
        elif output_format == 'pptx':
            prs = aligner.align_to_pptx(template_path)
            if prs:
                output_path = os.path.join(app.config['OUTPUT_FOLDER'], f"{output_filename}.pptx")
                prs.save(output_path)
                return send_file(output_path, as_attachment=True,
                               download_name=f"ClimateLenz_Report_{timestamp}.pptx")
            else:
                return jsonify({'error': 'python-pptx not installed'}), 500
        
        else:
            return jsonify({'error': 'Unsupported format'}), 400
    
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/preview', methods=['POST'])
def preview():
    """Return a preview of how material will be structured."""
    material_text = request.form.get('material', '').strip()
    if not material_text:
        return jsonify({'error': 'No material provided'}), 400
    
    processor = MaterialProcessor(material_text)
    
    preview_data = {
        'total_sections': len(processor.sections),
        'total_lines': sum(len(s['content']) for s in processor.sections),
        'sections': [
            {
                'title': s['title'],
                'line_count': len(s['content']),
                'preview': ' '.join(s['content'][:3]) + ('...' if len(s['content']) > 3 else '')
            }
            for s in processor.sections
        ]
    }
    
    return jsonify(preview_data)


@app.route('/download-all', methods=['POST'])
def download_all():
    """Download all three formats as a zip."""
    material_text = request.form.get('material', '').strip()
    if not material_text:
        return jsonify({'error': 'No material provided'}), 400
    
    processor = MaterialProcessor(material_text)
    aligner = TemplateAligner(processor)
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    zip_filename = f"climatelenz_all_formats_{timestamp}.zip"
    zip_path = os.path.join(app.config['OUTPUT_FOLDER'], zip_filename)
    
    with zipfile.ZipFile(zip_path, 'w') as zf:
        # Word
        if DOCX_AVAILABLE:
            doc = aligner.align_to_word()
            doc_path = os.path.join(app.config['OUTPUT_FOLDER'], f"temp_{timestamp}.docx")
            doc.save(doc_path)
            zf.write(doc_path, f"ClimateLenz_Report_{timestamp}.docx")
            os.remove(doc_path)
        
        # Excel
        if EXCEL_AVAILABLE:
            wb = aligner.align_to_excel()
            xlsx_path = os.path.join(app.config['OUTPUT_FOLDER'], f"temp_{timestamp}.xlsx")
            wb.save(xlsx_path)
            zf.write(xlsx_path, f"ClimateLenz_Report_{timestamp}.xlsx")
            os.remove(xlsx_path)
        
        # PowerPoint
        if PPTX_AVAILABLE:
            prs = aligner.align_to_pptx()
            pptx_path = os.path.join(app.config['OUTPUT_FOLDER'], f"temp_{timestamp}.pptx")
            prs.save(pptx_path)
            zf.write(pptx_path, f"ClimateLenz_Report_{timestamp}.pptx")
            os.remove(pptx_path)
    
    return send_file(zip_path, as_attachment=True, download_name=zip_filename)


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)